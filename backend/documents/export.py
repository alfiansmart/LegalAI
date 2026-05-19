"""Render a document to .docx / .md / .pdf-bytes.

Phase-2 ships .docx and .md. .pdf is approximated as .docx for now —
production export uses a headless renderer (Phase 3+).
"""
from __future__ import annotations

import io
import re

from docx import Document as DocxDocument
from docx.shared import Pt


def to_docx_bytes(markdown_text: str, title: str | None = None) -> bytes:
    """Convert minimal markdown (#, ##, ###, bold, lists, paragraphs) to .docx."""
    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    if title:
        doc.core_properties.title = title

    for raw in markdown_text.splitlines():
        line = raw.rstrip()
        if not line:
            doc.add_paragraph("")
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif re.match(r"^\s*[-*]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*[-*]\s+", "", line), style="List Bullet")
        elif re.match(r"^\s*\d+\.\s+", line):
            doc.add_paragraph(re.sub(r"^\s*\d+\.\s+", "", line), style="List Number")
        else:
            p = doc.add_paragraph()
            _apply_inline(p, line)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_markdown_bytes(markdown_text: str) -> bytes:
    return markdown_text.encode("utf-8")


_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_ITAL_RE = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")


def _apply_inline(paragraph, text: str) -> None:
    cursor = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > cursor:
            paragraph.add_run(text[cursor : m.start()])
        run = paragraph.add_run(m.group(1))
        run.bold = True
        cursor = m.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])
